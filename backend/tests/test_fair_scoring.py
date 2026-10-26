"""QA check for NFR-04 (DP-24): scores don't change when the candidate does.

ats_score, hiring_chance and match_keywords only ever take the resume's
extracted text and its format problems (see their docstrings) - never the
candidate's name, photo, gender, age or nationality. These tests prove that
end to end: two resumes that differ only in the candidate's name and/or
photo must come back with identical scores.
"""

from docx import Document
from PIL import Image

JOB = "Data analyst with SQL, Python and Tableau. SQL is a must. Kubernetes is a plus."

# Everything after the name/photo is identical across every resume built below.
BODY = [
    "jane@example.com | +998 90 123 45 67",
    "Summary",
    "Data analyst with five years in retail.",
    "Experience",
    "Data analyst, Acme, 2019 - 2024",
    "Cut report time by 40%",
    "Built dashboards for 120 managers",
    "Saved $50,000 a year",
    "Education",
    "Webster University",
    "Skills",
    "SQL, Python, Tableau, Excel, Power BI, Statistics, Pandas, Git",
]

# Names that differ by apparent gender and nationality. If any score changed
# across these, it would be reacting to who the candidate is, not what they wrote.
NAMES = ["John Smith", "Amina Yusupova", "Chen Wei", "Fatima Al-Sayed"]

PHOTO_COLORS = ["navy", "crimson", "gold"]


def _photo(tmp_path, label, color):
    path = tmp_path / f"{label}.png"
    Image.new("RGB", (40, 40), color).save(path)
    return path


def _build_docx(tmp_path, name, photo_color=None):
    document = Document()
    document.add_paragraph(name)
    if photo_color is not None:
        document.add_picture(str(_photo(tmp_path, name.replace(" ", "_"), photo_color)))
    for line in BODY:
        document.add_paragraph(line)
    path = tmp_path / f"{name.replace(' ', '_')}_{photo_color}.docx"
    document.save(str(path))
    return path


def _upload(client, path):
    response = client.post("/api/resumes", files={"file": (path.name, path.read_bytes())})
    assert response.status_code == 201
    return response.json()["id"]


def _all_scores(client, resume_id):
    ats = client.get(f"/api/resumes/{resume_id}/ats-score")
    assert ats.status_code == 200
    keywords = client.post(
        f"/api/resumes/{resume_id}/keyword-match", json={"job_description": JOB}
    )
    assert keywords.status_code == 200
    hiring = client.post(
        f"/api/resumes/{resume_id}/hiring-chance", json={"job_description": JOB}
    )
    assert hiring.status_code == 200
    return ats.json(), keywords.json(), hiring.json()


def _assert_same_scores(a, b):
    ats_a, keywords_a, hiring_a = a
    ats_b, keywords_b, hiring_b = b
    assert ats_a["score"] == ats_b["score"]
    assert ats_a["deductions"] == ats_b["deductions"]
    assert keywords_a["percentage"] == keywords_b["percentage"]
    assert hiring_a["percentage"] == hiring_b["percentage"]
    assert hiring_a["level"] == hiring_b["level"]
    assert [f["score"] for f in hiring_a["factors"]] == [f["score"] for f in hiring_b["factors"]]


def test_changing_the_candidates_name_does_not_change_any_score(client, tmp_path):
    baseline_id = _upload(client, _build_docx(tmp_path, NAMES[0]))
    baseline = _all_scores(client, baseline_id)

    for name in NAMES[1:]:
        resume_id = _upload(client, _build_docx(tmp_path, name))
        _assert_same_scores(baseline, _all_scores(client, resume_id))


def test_changing_the_candidates_photo_does_not_change_any_score(client, tmp_path):
    baseline_id = _upload(client, _build_docx(tmp_path, NAMES[0], photo_color=PHOTO_COLORS[0]))
    baseline = _all_scores(client, baseline_id)

    for color in PHOTO_COLORS[1:]:
        resume_id = _upload(client, _build_docx(tmp_path, NAMES[0], photo_color=color))
        _assert_same_scores(baseline, _all_scores(client, resume_id))


def test_changing_both_name_and_photo_together_does_not_change_any_score(client, tmp_path):
    baseline_id = _upload(client, _build_docx(tmp_path, NAMES[0], photo_color=PHOTO_COLORS[0]))
    baseline = _all_scores(client, baseline_id)

    for name, color in zip(NAMES[1:], PHOTO_COLORS[1:] * 2, strict=False):
        resume_id = _upload(client, _build_docx(tmp_path, name, photo_color=color))
        _assert_same_scores(baseline, _all_scores(client, resume_id))
