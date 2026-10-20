from docx import Document

from app.format_checks import NO_PHONE, TABLES


def upload_docx(client, tmp_path, lines, table=False):
    path = tmp_path / "resume.docx"
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    if table:
        document.add_table(rows=1, cols=2).cell(0, 0).text = "Acme analyst"
    document.save(str(path))
    response = client.post("/api/resumes", files={"file": (path.name, path.read_bytes())})
    assert response.status_code == 201
    return response.json()["id"]


CLEAN_RESUME = [
    "Jane Doe",
    "jane@example.com",
    "+998 90 123 45 67",
    "Experience",
    "Data analyst, Acme",
    "Education",
    "Webster University",
    "Skills",
    "SQL",
]


def test_a_clean_resume_scores_100(client, tmp_path):
    resume_id = upload_docx(client, tmp_path, CLEAN_RESUME)

    response = client.get(f"/api/resumes/{resume_id}/ats-score")

    assert response.status_code == 200
    assert response.json() == {
        "id": resume_id,
        "score": 100,
        "max_score": 100,
        "deductions": [],
    }


def test_the_score_lists_what_each_problem_cost(client, tmp_path):
    lines = [line for line in CLEAN_RESUME if not line.startswith("+998")]
    resume_id = upload_docx(client, tmp_path, lines, table=True)

    body = client.get(f"/api/resumes/{resume_id}/ats-score").json()

    assert body["score"] == 100 - 20 - 5
    assert body["deductions"] == [
        {
            "code": "tables",
            "title": "Tables",
            "explanation": TABLES.explanation,
            "points": 20,
        },
        {
            "code": "no_phone",
            "title": "No phone number",
            "explanation": NO_PHONE.explanation,
            "points": 5,
        },
    ]


def test_the_same_resume_gets_the_same_score_every_time(client, tmp_path):
    resume_id = upload_docx(client, tmp_path, CLEAN_RESUME[:4], table=True)

    bodies = [client.get(f"/api/resumes/{resume_id}/ats-score").json() for _ in range(3)]

    assert bodies[0] == bodies[1] == bodies[2]


def test_the_score_of_an_unknown_resume_is_not_found(client):
    response = client.get(f"/api/resumes/{'0' * 32}/ats-score")

    assert response.status_code == 404
