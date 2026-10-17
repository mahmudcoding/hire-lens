from docx import Document


def upload_docx(client, tmp_path, build):
    path = tmp_path / "resume.docx"
    document = Document()
    build(document)
    document.save(str(path))
    response = client.post("/api/resumes", files={"file": (path.name, path.read_bytes())})
    assert response.status_code == 201
    return response.json()["id"]


def resume_with_a_table(document):
    for line in ["Jane Doe", "jane@example.com", "+998 90 123 45 67", "Experience"]:
        document.add_paragraph(line)
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Acme analyst"
    for line in ["Education", "Webster University", "Skills", "SQL"]:
        document.add_paragraph(line)


def test_format_problems_of_an_uploaded_resume_are_returned(client, tmp_path):
    resume_id = upload_docx(client, tmp_path, resume_with_a_table)

    response = client.get(f"/api/resumes/{resume_id}/format-checks")

    assert response.status_code == 200
    body = response.json()
    assert body["id"] == resume_id
    assert [problem["code"] for problem in body["problems"]] == ["tables"]
    assert body["problems"][0]["title"] == "Tables"
    assert "table" in body["problems"][0]["explanation"]


def test_format_checks_of_an_unknown_resume_are_not_found(client):
    response = client.get(f"/api/resumes/{'0' * 32}/format-checks")

    assert response.status_code == 404
