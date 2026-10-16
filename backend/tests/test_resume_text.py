from docx import Document
from PIL import Image


def upload(client, path):
    with open(path, "rb") as file:
        response = client.post("/api/resumes", files={"file": (path.name, file.read())})
    assert response.status_code == 201
    return response.json()["id"]


def test_text_of_an_uploaded_docx_is_returned(client, tmp_path):
    path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Experience: Data analyst at Acme")
    document.save(str(path))
    resume_id = upload(client, path)

    response = client.get(f"/api/resumes/{resume_id}/text")

    assert response.status_code == 200
    assert response.json() == {
        "id": resume_id,
        "text": "Jane Doe\nExperience: Data analyst at Acme",
    }


def test_a_scanned_pdf_gets_a_clear_error_not_a_crash(client, tmp_path):
    path = tmp_path / "scan.pdf"
    Image.new("RGB", (600, 800), "white").save(path, "PDF")
    resume_id = upload(client, path)

    response = client.get(f"/api/resumes/{resume_id}/text")

    assert response.status_code == 422
    assert "probably a scanned image" in response.json()["detail"]


def test_an_unknown_resume_is_not_found(client):
    response = client.get(f"/api/resumes/{'0' * 32}/text")

    assert response.status_code == 404
    assert response.json() == {"detail": "We couldn't find that resume. Please upload it again."}
