from docx import Document

from app import format_checks
from app.ats_score import MAX_SCORE, WEIGHTS, calculate_ats_score
from app.extraction import extract_text
from app.format_checks import (
    COLUMNS,
    IMAGES,
    NO_PHONE,
    TABLES,
    FormatProblem,
    check_format,
)


def problem(code):
    return FormatProblem(code=code, title=code.title(), explanation="")


def score_docx(path):
    return calculate_ats_score(check_format(path, "docx", extract_text(path, "docx")))


def test_a_resume_with_no_problems_scores_100():
    result = calculate_ats_score([])

    assert result.score == MAX_SCORE
    assert result.deductions == []


def test_each_problem_takes_off_its_weight():
    result = calculate_ats_score([TABLES, NO_PHONE])

    assert result.score == 100 - 20 - 5
    assert [(d.code, d.points) for d in result.deductions] == [("tables", 20), ("no_phone", 5)]


def test_deductions_are_listed_biggest_first_with_their_titles():
    result = calculate_ats_score([NO_PHONE, IMAGES, COLUMNS])

    assert [d.code for d in result.deductions] == ["columns", "images", "no_phone"]
    assert result.deductions[0].title == "Multiple columns"


def test_the_score_never_goes_below_0():
    result = calculate_ats_score([problem(code) for code in WEIGHTS])

    assert sum(WEIGHTS.values()) > MAX_SCORE
    assert result.score == 0
    assert len(result.deductions) == len(WEIGHTS)


def test_the_score_stays_between_0_and_100_for_any_set_of_problems():
    codes = list(WEIGHTS)
    for mask in range(2 ** len(codes)):
        problems = [problem(code) for i, code in enumerate(codes) if mask >> i & 1]
        assert 0 <= calculate_ats_score(problems).score <= MAX_SCORE


def test_every_format_problem_has_a_weight():
    codes = {
        value.code
        for value in vars(format_checks).values()
        if isinstance(value, FormatProblem)
    }
    codes.add(format_checks._missing_headings(["skills"]).code)

    assert codes == set(WEIGHTS)
    assert all(points > 0 for points in WEIGHTS.values())


def test_the_same_resume_always_gets_the_same_score(tmp_path):
    path = tmp_path / "resume.docx"
    document = Document()
    for line in ["Jane Doe", "jane@example.com", "Experience", "Data analyst, Acme"]:
        document.add_paragraph(line)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "SQL"
    document.save(str(path))

    results = [score_docx(path) for _ in range(3)]

    assert results[0] == results[1] == results[2]
    # A table, missing education and skills headings, and no phone number.
    assert results[0].score == 100 - 20 - 15 - 5

