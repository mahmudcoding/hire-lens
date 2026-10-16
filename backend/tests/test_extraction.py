import io
import zipfile

import pytest
from docx import Document
from docx.oxml import parse_xml
from fpdf import FPDF
from PIL import Image

from app.extraction import UnreadableResume, extract_text


def make_pdf(path, pages, user_password=None):
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    if user_password:
        pdf.set_encryption(owner_password="owner", user_password=user_password)
    for text in pages:
        pdf.add_page()
        pdf.multi_cell(0, 8, text)
    pdf.output(str(path))
    return path


def make_docx(path, build):
    document = Document()
    build(document)
    document.save(str(path))
    return path


TEXT_BOX = """
<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
     xmlns:v="urn:schemas-microsoft-com:vml">
  <w:r>
    <mc:AlternateContent>
      <mc:Choice Requires="wps">
        <w:drawing><wps:txbx><w:txbxContent>
          <w:p><w:r><w:t>Contact: jane@example.com</w:t></w:r></w:p>
        </w:txbxContent></wps:txbx></w:drawing>
      </mc:Choice>
      <mc:Fallback>
        <w:pict><v:textbox><w:txbxContent>
          <w:p><w:r><w:t>Contact: jane@example.com</w:t></w:r></w:p>
        </w:txbxContent></v:textbox></w:pict>
      </mc:Fallback>
    </mc:AlternateContent>
  </w:r>
</w:p>
"""


def test_pdf_text_is_extracted(tmp_path):
    path = make_pdf(
        tmp_path / "resume.pdf",
        ["Jane Doe\njane@example.com\n\nExperience\nData analyst at Acme"],
    )

    text = extract_text(path, "pdf")

    assert "Jane Doe" in text
    assert "jane@example.com" in text
    assert "Data analyst at Acme" in text


def test_every_page_of_a_pdf_is_read_in_order(tmp_path):
    path = make_pdf(tmp_path / "resume.pdf", ["Page one: Experience", "Page two: Education"])

    text = extract_text(path, "pdf")

    assert text.index("Page one: Experience") < text.index("Page two: Education")


def test_a_scanned_pdf_gets_a_clear_message(tmp_path):
    path = tmp_path / "scan.pdf"
    Image.new("RGB", (600, 800), "white").save(path, "PDF")

    with pytest.raises(UnreadableResume, match="probably a scanned image"):
        extract_text(path, "pdf")


def test_a_damaged_pdf_gets_a_clear_message(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.7\nthis is not really a pdf")

    with pytest.raises(UnreadableResume, match="may be damaged or password-protected"):
        extract_text(path, "pdf")


def test_a_password_protected_pdf_gets_a_clear_message(tmp_path):
    path = make_pdf(tmp_path / "locked.pdf", ["Jane Doe"], user_password="secret")

    with pytest.raises(UnreadableResume, match="may be damaged or password-protected"):
        extract_text(path, "pdf")


def test_docx_paragraphs_and_tables_are_read_in_order(tmp_path):
    def build(document):
        document.add_paragraph("Jane Doe")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Skills"
        table.cell(0, 1).text = "Python, SQL"
        document.add_paragraph("Education")

    text = extract_text(make_docx(tmp_path / "resume.docx", build), "docx")

    assert text.split("\n") == ["Jane Doe", "Skills", "Python, SQL", "Education"]


def test_docx_runs_in_one_paragraph_stay_on_one_line(tmp_path):
    def build(document):
        paragraph = document.add_paragraph("Data ")
        paragraph.add_run("analyst").bold = True
        paragraph.add_run(" at Acme")

    text = extract_text(make_docx(tmp_path / "resume.docx", build), "docx")

    assert text == "Data analyst at Acme"


def test_docx_text_box_text_is_read_once(tmp_path):
    def build(document):
        heading = document.add_paragraph("Jane Doe")
        heading._p.addnext(parse_xml(TEXT_BOX))

    text = extract_text(make_docx(tmp_path / "resume.docx", build), "docx")

    assert text == "Jane Doe\nContact: jane@example.com"


def test_docx_header_is_left_out_like_an_ats_would(tmp_path):
    def build(document):
        document.sections[0].header.paragraphs[0].text = "Jane Doe, jane@example.com"
        document.add_paragraph("Experience")

    text = extract_text(make_docx(tmp_path / "resume.docx", build), "docx")

    assert text == "Experience"


def test_a_docx_without_text_gets_a_clear_message(tmp_path):
    path = make_docx(tmp_path / "empty.docx", lambda document: None)

    with pytest.raises(UnreadableResume, match="couldn't find any text"):
        extract_text(path, "docx")


def test_a_damaged_docx_gets_a_clear_message(tmp_path):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "not xml")
    path = tmp_path / "broken.docx"
    path.write_bytes(buffer.getvalue())

    with pytest.raises(UnreadableResume, match="may be damaged"):
        extract_text(path, "docx")
