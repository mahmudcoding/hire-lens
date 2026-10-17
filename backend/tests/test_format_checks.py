from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from fpdf import FPDF
from PIL import Image

from app.extraction import extract_text
from app.format_checks import check_format

TEXT_BOX = """
<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <w:r><w:drawing><wps:txbx><w:txbxContent>
    <w:p><w:r><w:t>jane@example.com</w:t></w:r></w:p>
  </w:txbxContent></wps:txbx></w:drawing></w:r>
</w:p>
"""

PLAIN_LINES = [
    "Jane Doe",
    "jane@example.com | +998 90 123 45 67",
    "Summary",
    "Data analyst with four years of experience in retail reporting.",
    "Experience",
    "Data analyst, Acme (2021 - 2025)",
    "Education",
    "BSc Computer Science, Webster University",
    "Skills",
    "Python, SQL, Power BI",
]


def check(path):
    file_type = path.suffix.removeprefix(".")
    return check_format(path, file_type, extract_text(path, file_type))


def codes(path):
    return [problem.code for problem in check(path)]


def new_pdf():
    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    pdf.add_page()
    return pdf


def write_lines(pdf, x, lines, width=100, top=20, spacing=8):
    for index, line in enumerate(lines):
        pdf.set_xy(x, top + index * spacing)
        pdf.cell(width, 8, line)


def save_pdf(pdf, tmp_path):
    path = tmp_path / "resume.pdf"
    pdf.output(str(path))
    return path


def save_docx(document, tmp_path):
    path = tmp_path / "resume.docx"
    document.save(str(path))
    return path


def plain_docx():
    document = Document()
    for line in PLAIN_LINES:
        document.add_paragraph(line)
    return document


def picture(tmp_path):
    path = tmp_path / "photo.png"
    Image.new("RGB", (40, 40), "navy").save(path)
    return path


def test_a_plain_pdf_has_no_layout_problems(tmp_path):
    pdf = new_pdf()
    write_lines(pdf, 10, PLAIN_LINES, width=190)

    assert check(save_pdf(pdf, tmp_path)) == []


def test_a_table_in_a_pdf_is_found(tmp_path):
    pdf = new_pdf()
    with pdf.table() as table:
        for cells in [("Company", "Years"), ("Acme", "2021 - 2025"), ("Beta", "2019 - 2021")]:
            row = table.row()
            for cell in cells:
                row.cell(cell)

    assert "tables" in codes(save_pdf(pdf, tmp_path))


def test_a_two_column_pdf_is_found(tmp_path):
    pdf = new_pdf()
    write_lines(pdf, 10, ["Jane Doe", "jane@example.com", "Skills", "Python", "SQL", "Power BI"], 60)
    write_lines(
        pdf,
        90,
        [
            "Experience",
            "Data analyst, Acme Corporation",
            "Built weekly sales dashboards",
            "Education",
            "BSc Computer Science",
            "Webster University",
        ],
        top=24,
        spacing=7,
    )

    assert "columns" in codes(save_pdf(pdf, tmp_path))


def test_right_aligned_dates_and_cities_are_not_a_second_column(tmp_path):
    pdf = new_pdf()
    for index in range(6):
        top = 20 + index * 24
        write_lines(pdf, 10, ["Data analyst at Acme Corporation"], top=top)
        write_lines(pdf, 150, ["Tashkent, Uzbekistan" if index % 2 else "2021 - 2025"], top=top)
        write_lines(pdf, 10, ["- Built short weekly reports"], top=top + 8)

    assert "columns" not in codes(save_pdf(pdf, tmp_path))


def test_an_image_in_a_pdf_is_found(tmp_path):
    pdf = new_pdf()
    write_lines(pdf, 10, PLAIN_LINES, width=190)
    pdf.image(str(picture(tmp_path)), x=170, y=10, w=20)

    assert codes(save_pdf(pdf, tmp_path)) == ["images"]


def test_a_plain_docx_has_no_layout_problems(tmp_path):
    assert check(save_docx(plain_docx(), tmp_path)) == []


def test_a_table_in_a_docx_is_found(tmp_path):
    document = plain_docx()
    document.add_table(rows=2, cols=2).cell(0, 0).text = "Acme"

    assert codes(save_docx(document, tmp_path)) == ["tables"]


def test_columns_in_a_docx_are_found(tmp_path):
    document = plain_docx()
    cols = document.sections[0]._sectPr.makeelement(qn("w:cols"), {qn("w:num"): "2"})
    document.sections[0]._sectPr.append(cols)

    assert codes(save_docx(document, tmp_path)) == ["columns"]


def test_text_boxes_in_a_docx_are_found(tmp_path):
    document = plain_docx()
    document.element.body.insert(0, parse_xml(TEXT_BOX))

    assert codes(save_docx(document, tmp_path)) == ["text_boxes"]


def test_a_picture_in_a_docx_is_found(tmp_path):
    document = plain_docx()
    document.add_picture(str(picture(tmp_path)))

    assert codes(save_docx(document, tmp_path)) == ["images"]


def test_a_logo_in_a_docx_header_is_found(tmp_path):
    document = plain_docx()
    header = document.sections[0].header
    header.paragraphs[0].add_run().add_picture(str(picture(tmp_path)))

    assert codes(save_docx(document, tmp_path)) == ["images"]


def test_unreadable_font_text_is_reported(tmp_path):
    path = save_docx(plain_docx(), tmp_path)

    problems = check_format(path, "docx", extract_text(path, "docx") + "\n(cid:12)(cid:15)")

    assert [problem.code for problem in problems] == ["unreadable_font"]


def test_missing_standard_headings_are_named(tmp_path):
    document = Document()
    for line in ["Jane Doe", "jane@example.com", "+998 90 123 45 67", "Work", "Acme analyst"]:
        document.add_paragraph(line)

    [problem] = check(save_docx(document, tmp_path))

    assert problem.code == "missing_headings"
    assert "experience, education and skills" in problem.explanation


def test_missing_email_and_phone_are_reported(tmp_path):
    document = Document()
    for line in PLAIN_LINES[:1] + PLAIN_LINES[2:]:
        document.add_paragraph(line)

    assert codes(save_docx(document, tmp_path)) == ["no_email", "no_phone"]


def test_contact_details_only_in_the_header_are_reported(tmp_path):
    document = Document()
    for line in PLAIN_LINES[:1] + PLAIN_LINES[2:]:
        document.add_paragraph(line)
    document.sections[0].header.paragraphs[0].text = "jane@example.com | +998 90 123 45 67"

    assert codes(save_docx(document, tmp_path)) == ["contact_in_header"]


def test_a_phone_missing_everywhere_is_reported_with_a_header_email(tmp_path):
    document = Document()
    for line in PLAIN_LINES[:1] + PLAIN_LINES[2:]:
        document.add_paragraph(line)
    document.sections[0].header.paragraphs[0].text = "jane@example.com"

    assert codes(save_docx(document, tmp_path)) == ["contact_in_header", "no_phone"]


def test_every_problem_has_a_title_and_an_explanation(tmp_path):
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Acme analyst"
    document.add_picture(str(picture(tmp_path)))

    problems = check(save_docx(document, tmp_path))

    assert len(problems) >= 4
    assert all(problem.title and problem.explanation for problem in problems)
